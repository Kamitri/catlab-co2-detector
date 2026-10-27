"""End-to-end test of built application with simulator, tauri-driver and virtual display (Linux only)."""

import os
import signal
import socket
import sqlite3
import subprocess
import sys
import time
from collections.abc import Callable
from pathlib import Path
from typing import Annotated

import openpyxl
import typer
from selenium import webdriver
from selenium.common.exceptions import WebDriverException
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.common.actions.action_builder import ActionBuilder
from selenium.webdriver.common.actions.mouse_button import MouseButton
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from Xlib import X, display, protocol

ROOT = Path(__file__).resolve().parents[1]
APPLICATION = ROOT / "app/src-tauri/target/debug/CATlab CO2 Monitor"
DATA = ROOT / "CO2 NDIR Monitor/KQ/Na-ZY (09-11).txt"


def start(command: list[str], **kwargs) -> subprocess.Popen:
    """Start background process."""
    return subprocess.Popen(command, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, **kwargs)


def session() -> webdriver.Remote:
    """Open WebDriver session on application."""
    options = webdriver.ChromeOptions()
    options.set_capability("browserName", "wry")
    options.set_capability("tauri:options", {"application": str(APPLICATION)})
    driver = webdriver.Remote("http://127.0.0.1:4444", options=options)
    driver.implicitly_wait(2)
    return driver


def button(driver: webdriver.Remote, text: str):
    """Find visible button by label."""
    return driver.find_element(By.XPATH, f"//button[normalize-space()='{text}']")


def wait_until(condition: Callable[[], bool], timeout: float = 30, message: str = "Condition not reached") -> None:
    """Poll condition until true; raise after timeout."""
    deadline = time.monotonic() + timeout
    while not condition():
        if time.monotonic() > deadline:
            raise TimeoutError(message)
        time.sleep(0.05)


def page_text(driver: webdriver.Remote) -> str:
    """Return visible text of page."""
    return driver.find_element(By.TAG_NAME, "body").text


def wait_text(driver: webdriver.Remote, text: str, timeout: float = 30, count: int = 1) -> None:
    """Wait until text appears on page at least count times."""
    wait_until(lambda: page_text(driver).count(text) >= count, timeout, f"Text not shown: {text}")


def co2_value(driver: webdriver.Remote) -> float:
    """Return current CO₂ value of status panel, or 0 before first measurement."""
    text = driver.find_element(By.CSS_SELECTOR, "aside span.text-4xl").text.replace(",", "")
    return float(text) if text.replace(".", "").isdigit() else 0.0


def close_window(title: str) -> None:
    """Send WM_DELETE_WINDOW to window, as title bar close button does."""
    result = subprocess.run(
        ["xdotool", "search", "--name", title], env=dict(os.environ, DISPLAY=":99"), capture_output=True, check=True, text=True
    )
    connection = display.Display(":99")
    delete = connection.intern_atom("WM_DELETE_WINDOW")
    protocols = connection.intern_atom("WM_PROTOCOLS")
    for window_id in result.stdout.split():
        window = connection.create_resource_object("window", int(window_id))
        event = protocol.event.ClientMessage(window=window, client_type=protocols, data=(32, [delete, X.CurrentTime, 0, 0, 0]))
        window.send_event(event)
    connection.flush()


def window_visible(title: str) -> bool:
    """Return whether window with title is mapped."""
    result = subprocess.run(
        ["xdotool", "search", "--onlyvisible", "--name", title], env=dict(os.environ, DISPLAY=":99"), capture_output=True, check=False
    )
    return bool(result.stdout.strip())


SOAP_TIMES = ["2.62", "2.75", "2.81", "2.68", "2.9", "2.71", "2.77", "2.66", "2.85", "2.73"]


def dialog_button(driver: webdriver.Remote, text: str):
    """Find button by label in open dialog."""
    return driver.find_element(By.XPATH, f"//div[@role='dialog' or @role='alertdialog']//button[normalize-space()='{text}']")


def check_export(work: Path) -> None:
    """Check sheet of first export: name date_sample, soap film times, and flow rate from times."""
    path = next((work / "exports").glob("*Na-ZY*.xlsx"))
    sheet = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    assert sheet.title.endswith("_Na-ZY"), sheet.title
    times = [sheet.cell(row, 17).value for row in range(8, 18)]
    assert times == [float(time) for time in SOAP_TIMES], times
    flow = next(sheet.cell(row, 12).value for row in range(1, 60) if sheet.cell(row, 11).value == "Flow rate")
    expected = sum(300 / float(time) for time in SOAP_TIMES) / len(SOAP_TIMES)
    assert abs(flow - expected) < 1e-9, (flow, expected)
    formats = {sheet.cell(row, 11).value: sheet.cell(row, 12).number_format for row in range(1, 60)}
    assert formats["Mass of sorbent"] == "0.0000" and formats["Soap film volume"] == "0.0", formats
    print(f"Export: sheet {sheet.title}, flow rate {flow:.2f} mL/min from {len(times)} soap film times")


def query_count(work: Path, sql: str) -> int:
    """Return count from query on application database."""
    with sqlite3.connect(work / "data/com.catlab.co2monitor/co2-monitor.db") as connection:
        return connection.execute(sql).fetchone()[0]


def measurement_count(work: Path) -> int:
    """Count stored measurements in application database."""
    return query_count(work, "SELECT COUNT(*) FROM measurements")


IN_RUN = "EXISTS (SELECT 1 FROM runs WHERE timestamp >= started_at AND timestamp <= COALESCE(finished_at, timestamp))"


def axis_ranges(driver: webdriver.Remote) -> list[list[float]]:
    """Return rounded x and y axis ranges of chart."""
    ranges = driver.execute_script(
        "const plot = document.querySelector('.js-plotly-plot'); return [plot.layout.xaxis.range, plot.layout.yaxis.range]"
    )
    return [[round(value, 1) for value in axis] for axis in ranges]


def port_open(port: int) -> bool:
    """Return whether local TCP port accepts connections."""
    with socket.socket() as connection:
        return connection.connect_ex(("127.0.0.1", port)) == 0


def fill(driver: webdriver.Remote, element_id: str, value: str) -> None:
    """Replace input value."""
    element = driver.find_element(By.ID, element_id)
    element.send_keys(Keys.CONTROL, "a")
    element.send_keys(Keys.BACKSPACE)
    element.send_keys(value)


def main(
    work: Annotated[Path, typer.Option(help="Work folder")] = Path("/tmp/co2-e2e"),
    step: Annotated[int, typer.Option(help="Data points per poll")] = 20,
    poll_interval: Annotated[int, typer.Option(help="Poll interval of application (ms)")] = 50,
) -> None:
    """Run end-to-end scenario and save screenshots in work folder."""
    subprocess.run(["rm", "-rf", str(work)], check=True)
    (work / "shots").mkdir(parents=True)
    environment = dict(
        os.environ,
        DISPLAY=":99",
        CO2_EXTRA_PORTS=str(work / "app"),
        CO2_POLL_INTERVAL_MS=str(poll_interval),
        XDG_DATA_HOME=str(work / "data"),
        # Time zone other than UTC, so clock axis checks catch time zone errors
        TZ="Asia/Ho_Chi_Minh",
    )
    processes = [start(["Xvfb", ":99", "-screen", "0", "1440x900x24"])]
    wait_until(lambda: Path("/tmp/.X11-unix/X99").exists(), 10, "Virtual display not started")
    processes.append(start(["socat", f"pty,raw,echo=0,link={work / 'device'}", f"pty,raw,echo=0,link={work / 'app'}"]))
    wait_until(lambda: (work / "device").exists() and (work / "app").exists(), 10, "Virtual serial port not created")
    simulator = start([sys.executable, str(ROOT / "tools/simulator.py"), str(work / "device"), "--data", str(DATA), "--step", str(step)])
    processes.append(simulator)
    processes.append(start(["tauri-driver", "--native-driver", "/usr/bin/WebKitWebDriver"], env=environment))
    wait_until(lambda: port_open(4444), 10, "tauri-driver not started")
    started = time.monotonic()

    def shot(driver: webdriver.Remote, name: str) -> None:
        driver.save_screenshot(str(work / "shots" / f"{name}.png"))

    try:
        driver = session()
        driver.set_window_size(1440, 900)
        wait_text(driver, "Connected to")
        wait_until(lambda: co2_value(driver) > 0, message="Measurement not shown")
        # Pointer is not on chart: line and label at newest data point
        latest = "return document.querySelector('.js-plotly-plot').layout.annotations.some((note) => note.text.endsWith(' ppm'))"
        wait_until(lambda: driver.execute_script(latest), 5, "Newest data point not shown")
        shot(driver, "01-monitor")

        button(driver, "Settings").click()
        fill(driver, "folder", str(work / "exports"))
        fill(driver, "equilibriumBandPercent", "5")
        # Dark theme; later screenshots show it
        driver.find_element(By.ID, "theme").click()
        driver.find_element(By.XPATH, "//div[@role='option'][normalize-space()='Dark']").click()
        shot(driver, "02-settings")
        button(driver, "Save").click()
        wait_text(driver, "Settings saved")
        dark = "return document.documentElement.classList.contains('dark')"
        wait_until(lambda: driver.execute_script(dark), 5, "Dark theme not applied")

        # Sample is configured before Start; Start then starts run at once
        assert not button(driver, "Start").is_enabled(), "Start enabled without sample"
        button(driver, "Sample").click()
        fill(driver, "sample", "Na-ZY")
        fill(driver, "operator", "Test")
        fill(driver, "massG", "0.56")
        # Windows scaled to fast poll interval: data of 1 s real time covers about 400 s of run
        fill(driver, "coWindowS", "1")
        # Mass and soap film volume show decimal places of balance and flowmeter after field loses focus
        assert driver.find_element(By.ID, "massG").get_attribute("value") == "0.5600"
        assert driver.find_element(By.ID, "soapVolumeMl").get_attribute("value") == "5.0"
        shot(driver, "03-sample")
        dialog_button(driver, "Save").click()
        wait_until(lambda: button(driver, "Start").is_enabled(), 5, "Start not enabled after sample saved")
        button(driver, "Start").click()
        wait_text(driver, "Run started: Na-ZY", 2)
        wait_until(lambda: co2_value(driver) > 15000, message="Breakthrough not reached")
        shot(driver, "04-running")

        # Device stops replying for 2 s, then recovers
        frozen = time.monotonic()
        os.kill(simulator.pid, signal.SIGSTOP)
        wait_text(driver, "Disconnected from")
        wait_until(lambda: time.monotonic() - frozen > 2)
        os.kill(simulator.pid, signal.SIGCONT)
        wait_text(driver, "Connected to", count=2)

        wait_text(driver, "Equilibrium reached", 60)
        shot(driver, "05-equilibrium")
        button(driver, "Finish").click()
        for index, value in enumerate(SOAP_TIMES):
            fill(driver, f"soapTime-{index}", value)
        # Add creates row 11; its delete button removes it again
        dialog_button(driver, "Add").click()
        driver.find_element(By.ID, "soapTime-10")
        driver.find_element(By.CSS_SELECTOR, "button[aria-label='Delete time 11']").click()
        wait_until(lambda: not driver.find_elements(By.ID, "soapTime-10"), 5, "Soap film row not deleted")
        shot(driver, "06-finish")
        dialog_button(driver, "Finish").click()
        wait_text(driver, "Exported")
        shot(driver, "07-finished")
        check_export(work)

        driver.find_element(By.XPATH, "//button[normalize-space()='History']").click()
        driver.find_element(By.XPATH, "//td[normalize-space()='Na-ZY']").click()
        wait_text(driver, "Stoichiometric time")
        shot(driver, "08-history")
        driver.find_element(By.XPATH, "//button[normalize-space()='About']").click()
        wait_text(driver, "All rights reserved")
        shot(driver, "08b-about")

        # Close button during run: window hides to tray and recording continues
        driver.find_element(By.XPATH, "//button[normalize-space()='Live']").click()
        button(driver, "Sample").click()
        fill(driver, "sample", "Crash test")
        dialog_button(driver, "Save").click()
        wait_until(lambda: not driver.find_elements(By.XPATH, "//div[@role='dialog']"), 5, "Sample dialog not closed")
        button(driver, "Start").click()
        wait_text(driver, "Run started: Crash test")
        close_window("CATlab CO₂ Monitor")
        wait_until(lambda: not window_visible("CATlab CO₂ Monitor"), 5, "Window still visible after close during run")
        count = measurement_count(work)
        wait_until(lambda: measurement_count(work) >= count + 5, 5, "Recording stopped after window close")
        print("Close during run: window hidden, recording continues")

        # Crash during run: application is killed, then restarted
        subprocess.run(["pkill", "-9", "-f", str(APPLICATION)], check=False)
        wait_until(lambda: subprocess.run(["pgrep", "-f", str(APPLICATION)], capture_output=True, check=False).returncode != 0, 10)
        # Session of killed application can fail to close; new session follows
        try:
            driver.quit()
        except WebDriverException as error:
            print(f"Old session not closed: {error.msg}")
        driver = session()
        driver.set_window_size(1440, 900)
        wait_text(driver, "Run in progress")
        shot(driver, "09-recovery")
        button(driver, "Resume").click()
        wait_text(driver, "Connected to", count=3)
        shot(driver, "10-resumed")
        button(driver, "Abort").click()
        driver.find_element(By.XPATH, "//div[@role='alertdialog']//button[normalize-space()='Abort']").click()
        wait_text(driver, "Exported", count=2)
        shot(driver, "11-aborted")

        # Several runs to one workbook, one sheet per run (native save dialog is skipped)
        combined = work / "exports" / "combined.xlsx"
        driver.execute_async_script(
            "const done = arguments[arguments.length - 1];"
            "window.__TAURI_INTERNALS__.invoke('export_runs', {ids: [1, 2], path: arguments[0]}).then(done, done)",
            str(combined),
        )
        sheets = openpyxl.load_workbook(combined).sheetnames
        assert len(sheets) == 2 and sheets[0].endswith("_Na-ZY") and sheets[1].endswith("_Crash test"), sheets
        print("Export of several runs:", sheets)

        # Chart view checks run last: after WebDriver pointer actions, WebKit WebDriver drops Shift in typing
        # Default view: fixed axes that do not follow each new measurement
        default = [[0, 5000], [0, 40000]]
        assert axis_ranges(driver) == default, axis_ranges(driver)
        count = measurement_count(work)
        wait_until(lambda: measurement_count(work) >= count + 20, 10, "Measurements stopped")
        assert axis_ranges(driver) == default, axis_ranges(driver)

        # Axes tool sets exact ranges
        button(driver, "Axes").click()
        fill(driver, "axes-xTo", "1000")
        fill(driver, "axes-yTo", "10000")
        shot(driver, "01b-axes")
        button(driver, "Apply").click()
        wait_until(lambda: axis_ranges(driver) == [[0, 1000], [0, 10000]], 5, "Axes not applied")

        # Space with drag pans while Zoom tool stays selected. Key events come from script: WebDriver key
        # actions leave keyboard state that drops Shift in later typing
        plot = driver.find_element(By.CSS_SELECTOR, ".js-plotly-plot .nsewdrag")
        space = "window.dispatchEvent(new KeyboardEvent(arguments[0], {code: 'Space', key: ' ', bubbles: true}))"
        driver.execute_script(space, "keydown")
        ActionChains(driver).click_and_hold(plot).move_by_offset(-150, 0).release().perform()
        driver.execute_script(space, "keyup")
        wait_until(lambda: axis_ranges(driver)[0][0] > 0, 5, "Space drag did not pan")

        # Middle button drag pans
        before = axis_ranges(driver)[0][0]
        action = ActionBuilder(driver)
        action.pointer_action.move_to(plot).pointer_down(MouseButton.MIDDLE).move_by(-150, 0).pointer_up(MouseButton.MIDDLE)
        action.perform()
        wait_until(lambda: axis_ranges(driver)[0][0] > before, 5, "Middle button drag did not pan")
        assert not driver.find_elements(By.CSS_SELECTOR, ".notifier-note"), "Plotly tip shown"
        button(driver, "Reset view").click()
        wait_until(lambda: axis_ranges(driver) == default, 5, "Reset view did not restore default view")
        print("Chart view: default, Axes, Space pan, middle button pan, Reset view")

        # Clear chart asks first, then deletes measurements outside runs up to now
        run_measurements = query_count(work, f"SELECT COUNT(*) FROM measurements WHERE {IN_RUN}")
        clock_start = "return document.querySelector('.js-plotly-plot').layout.xaxis2.range[0]"
        before = driver.execute_script(clock_start)
        button(driver, "Clear chart").click()
        wait_text(driver, "This action will delete ALL your current chart data. This action is irreversible.")
        shot(driver, "12-clear")
        driver.find_element(By.XPATH, "//div[@role='alertdialog']//button[normalize-space()='Cancel']").click()
        assert query_count(work, f"SELECT COUNT(*) FROM measurements WHERE NOT {IN_RUN}") > 0
        button(driver, "Clear chart").click()
        driver.find_element(By.XPATH, "//div[@role='alertdialog']//button[normalize-space()='Clear chart']").click()
        wait_text(driver, "Chart cleared")
        cleared = query_count(work, "SELECT MAX(timestamp) FROM messages WHERE text = 'Chart cleared'")
        remaining = query_count(work, f"SELECT COUNT(*) FROM measurements WHERE timestamp < {cleared - 1000} AND NOT {IN_RUN}")
        assert remaining == 0, remaining
        assert query_count(work, f"SELECT COUNT(*) FROM measurements WHERE {IN_RUN}") == run_measurements
        # Chart starts again at time 0 = time of Clear chart; clock axis shows that time
        wait_until(lambda: driver.execute_script(clock_start) > before, 5, "Chart not cleared")
        cleared_at = query_count(work, "SELECT MAX(timestamp) FROM messages WHERE text = 'Chart cleared'")
        assert abs(driver.execute_script(clock_start) - cleared_at) < 5000, (driver.execute_script(clock_start), cleared_at)
        shot(driver, "13-cleared")
        print("Clear chart: confirmation, measurements outside runs deleted, run measurements kept")

        # Delete button of message removes it from log and database
        messages = query_count(work, "SELECT COUNT(*) FROM messages")
        # Newest message is in view; log follows end of list
        driver.find_elements(By.CSS_SELECTOR, "button[aria-label='Delete message']")[-1].click()
        wait_until(lambda: query_count(work, "SELECT COUNT(*) FROM messages") == messages - 1, 5, "Message not deleted")
        print("Message deleted")

        # Delete in History removes checked runs after confirmation
        driver.find_element(By.XPATH, "//button[normalize-space()='History']").click()
        row = driver.find_element(By.XPATH, "//tr[td[normalize-space()='Crash test']]")
        row.find_element(By.CSS_SELECTOR, "button[role=checkbox]").click()
        button(driver, "Delete").click()
        wait_text(driver, "This action will delete Crash test. This action is irreversible.")
        shot(driver, "14-delete")
        dialog_button(driver, "Delete").click()
        wait_until(lambda: query_count(work, "SELECT COUNT(*) FROM runs") == 1, 5, "Run not deleted")
        print("Run deleted in History")

        # Data editor: typing in cell replaces value; Save keeps measurements and stores edited data
        driver.find_element(By.XPATH, "//td[normalize-space()='Na-ZY']").click()
        button(driver, "Edit data").click()
        wait_until(lambda: driver.find_elements(By.CSS_SELECTOR, "[data-row]"), 5, "Data not shown")
        driver.find_element(By.CSS_SELECTOR, "[aria-label='Data']").send_keys("50000", Keys.ENTER)
        shot(driver, "16-data-editor")
        dialog_button(driver, "Save").click()
        edited_rows = "SELECT COUNT(*) FROM run_data"
        wait_until(lambda: query_count(work, edited_rows) > 0, 5, "Edited data not saved")
        first_value = query_count(work, "SELECT value_ppm FROM run_data WHERE position = 0")
        assert first_value == 50000, first_value
        edited_badge = "return document.body.textContent.includes('Edited')"
        wait_until(lambda: driver.execute_script(edited_badge), 5, "Edited badge not shown")
        # Restore original removes edited data after confirmation
        button(driver, "Edit data").click()
        dialog_button(driver, "Restore original").click()
        driver.find_element(By.XPATH, "//div[@role='alertdialog']//button[normalize-space()='Restore original']").click()
        wait_until(lambda: query_count(work, edited_rows) == 0, 5, "Edited data not removed")
        print("Data editor: edit, Save, Restore original")

        # Reset to defaults fills form with default values; Cancel keeps saved settings
        button(driver, "Settings").click()
        dialog_button(driver, "Reset to defaults").click()
        assert driver.find_element(By.ID, "equilibriumBandPercent").get_attribute("value") == "3"
        assert driver.find_element(By.ID, "massPrecisionG").get_attribute("value") == "0.0001"
        theme_text = driver.find_element(By.ID, "theme").get_attribute("textContent")
        assert theme_text == "Light", theme_text
        shot(driver, "15-settings-defaults")
        dialog_button(driver, "Cancel").click()
        assert driver.execute_script(dark), "Cancel changed theme"
        print("Settings: dark theme, Reset to defaults")
        driver.quit()
    finally:
        for process in reversed(processes):
            process.terminate()
        subprocess.run(["pkill", "-f", str(APPLICATION)], check=False)
    print(f"Scenario passed in {time.monotonic() - started:.0f} s")
    print("Exports:", sorted(path.name for path in (work / "exports").glob("*")))


if __name__ == "__main__":
    typer.run(main)
